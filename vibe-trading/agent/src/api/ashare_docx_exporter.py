import io
import re
import datetime
from typing import Dict, Any, List, Optional
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_shading(cell, color_hex: str):
    """设置单元格背景底色"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top: int = 100, bottom: int = 100, left: int = 140, right: int = 140):
    """设置单元格内部舒适呼吸边距（单位 dxa）"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>\n'
        f'  <w:top w:w="{top}" w:type="dxa"/>\n'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>\n'
        f'  <w:left w:w="{left}" w:type="dxa"/>\n'
        f'  <w:right w:w="{right}" w:type="dxa"/>\n'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1", sz="4"):
    """设置单元格边框线条"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{top}"/>\n'
        f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{left}"/>\n'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{bottom}"/>\n'
        f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{right}"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def set_run_fonts(run, east_asia="PingFang SC", ascii_font="Calibri", is_code=False):
    """设置中西文复合字体"""
    rPr = run._r.get_or_add_rPr()
    if is_code:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Consolas" w:hAnsi="Consolas" w:eastAsia="PingFang SC"/>')
    else:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{ascii_font}" w:hAnsi="{ascii_font}" w:eastAsia="{east_asia}"/>')
    rPr.append(rFonts)

def add_styled_paragraph(doc, text: str, font_size: float = 10.0, bold: bool = False, color_rgb: RGBColor = None, space_before: float = 4, space_after: float = 4, line_spacing: float = 1.35, align = None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    
    tokens = re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = p.add_run(token[2:-2])
            run.bold = True
            run.font.size = Pt(font_size)
            run.font.color.rgb = color_rgb or RGBColor(15, 23, 42)
            set_run_fonts(run)
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = p.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(font_size - 1)
            run.font.color.rgb = RGBColor(79, 70, 229)
            set_run_fonts(run, is_code=True)
        else:
            run = p.add_run(token)
            run.bold = bold
            run.font.size = Pt(font_size)
            if color_rgb:
                run.font.color.rgb = color_rgb
            else:
                run.font.color.rgb = RGBColor(30, 41, 59)
            set_run_fonts(run)
    return p

def add_callout_box(doc, title: str, items: List[str], callout_type: str = "suggestion"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.77)
    
    if callout_type == "suggestion":
        bg_color = "EEF2FF"
        bar_color = "4F46E5"
        subtle_border = "C7D2FE"
        icon = "💡 "
        title_color = RGBColor(30, 27, 75)
    elif callout_type == "warning":
        bg_color = "FFFBEB"
        bar_color = "F59E0B"
        subtle_border = "FDE68A"
        icon = "⚠️ "
        title_color = RGBColor(120, 53, 15)
    elif callout_type == "question":
        bg_color = "F0F9FF"
        bar_color = "0284C7"
        subtle_border = "BAE6FD"
        icon = "💬 "
        title_color = RGBColor(3, 105, 161)
    else:
        bg_color = "F8FAFC"
        bar_color = "6366F1"
        subtle_border = "E2E8F0"
        icon = "📌 "
        title_color = RGBColor(51, 65, 85)

    set_cell_shading(cell, bg_color)
    set_cell_margins(cell, top=130, bottom=130, left=160, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{subtle_border}"/>\n'
        f'  <w:left w:val="single" w:sz="28" w:space="0" w:color="{bar_color}"/>\n'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{subtle_border}"/>\n'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="{subtle_border}"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p_title = cell.paragraphs[0]
    p_title.paragraph_format.space_before = Pt(3)
    p_title.paragraph_format.space_after = Pt(3)
    run_icon = p_title.add_run(icon)
    run_icon.font.size = Pt(11)
    run_title = p_title.add_run(title)
    run_title.bold = True
    run_title.font.size = Pt(11)
    run_title.font.color.rgb = title_color
    set_run_fonts(run_title)
    
    for item in items:
        p_item = cell.add_paragraph()
        p_item.paragraph_format.space_before = Pt(2)
        p_item.paragraph_format.space_after = Pt(2)
        p_item.paragraph_format.line_spacing = 1.3
        
        tokens = re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', item)
        for token in tokens:
            if not token: continue
            if token.startswith('**') and token.endswith('**') and len(token) >= 4:
                run = p_item.add_run(token[2:-2])
                run.bold = True
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(15, 23, 42)
                set_run_fonts(run)
            elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
                run = p_item.add_run(token[1:-1])
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(79, 70, 229)
                set_run_fonts(run, is_code=True)
            else:
                run = p_item.add_run(token)
                run.font.size = Pt(9.5)
                if callout_type == "question":
                    run.font.color.rgb = RGBColor(15, 23, 42)
                    run.bold = True
                else:
                    run.font.color.rgb = RGBColor(30, 41, 59)
                set_run_fonts(run)
    
    doc.add_paragraph()

def build_ashare_docx(code: str, name: str, date_str: str, markdown_content: str, stats: Dict[str, Any] = None) -> io.BytesIO:
    doc = docx.Document()
    
    # 1. 页面布局：标准 A4，边距各 0.75 in (1.9cm)，可用正文宽 6.77 in
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    
    # 2. 页眉与页脚
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp_run = hp.add_run("A股多因子量化投研深度诊断研报 · Vibe-Trading AI")
    hp_run.font.size = Pt(8)
    hp_run.font.color.rgb = RGBColor(148, 163, 184)
    set_run_fonts(hp_run)
    hp_pPr = hp._p.get_or_add_pPr()
    hp_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="1" w:color="E2E8F0"/></w:pBdr>')
    hp_pPr.append(hp_border)
    
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp_run = fp.add_run("Ashare Vibe AI Quantitative Research System")
    fp_run.font.size = Pt(8)
    fp_run.font.color.rgb = RGBColor(148, 163, 184)
    set_run_fonts(fp_run)
    
    stats = stats or {}
    close_price = stats.get("close", "--")
    turnover_yi = stats.get("turnover", "--")
    industry = stats.get("industry", "--")
    export_time = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    
    # 3. 顶部机构级大标题与副标题
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_main = p_title.add_run("A 股多因子量化投研深度诊断报告")
    run_main.bold = True
    run_main.font.size = Pt(18)
    run_main.font.color.rgb = RGBColor(30, 27, 75)
    set_run_fonts(run_main, east_asia="PingFang SC", ascii_font="Calibri")
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(12)
    run_sub = p_sub.add_run("VIBE-TRADING QUANTITATIVE RESEARCH REPORT")
    run_sub.bold = True
    run_sub.font.size = Pt(9)
    run_sub.font.color.rgb = RGBColor(79, 70, 229)
    set_run_fonts(run_sub, ascii_font="Calibri")
    
    # 4. 标的行情元数据表格 (Meta Card)
    meta_table = doc.add_table(rows=3, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    meta_col_widths = [Inches(1.2), Inches(2.18), Inches(1.2), Inches(2.19)]
    meta_data = [
        [("标的证券", True), (f"{code} {name}", False), ("交易日期", True), (date_str, False)],
        [("收盘价格", True), (f"{close_price}", False), ("成交金额", True), (f"{turnover_yi}", False)],
        [("所属行业", True), (f"{industry}", False), ("报告时间", True), (export_time, False)],
    ]
    
    for r_idx, row_vals in enumerate(meta_data):
        for c_idx, (text, is_label) in enumerate(row_vals):
            cell = meta_table.cell(r_idx, c_idx)
            cell.width = meta_col_widths[c_idx]
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            set_cell_borders(cell, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1")
            if is_label:
                set_cell_shading(cell, "F1F5F9")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(text)
            run.font.size = Pt(9)
            if is_label:
                run.bold = True
                run.font.color.rgb = RGBColor(71, 85, 105)
            else:
                run.font.color.rgb = RGBColor(15, 23, 42)
            set_run_fonts(run)
                
    doc.add_paragraph() # 空白分隔
    
    # 5. 正文 Markdown 流式解析与原生 OpenXML 转换
    clean_md = re.sub(r"<think>[\s\S]*?</think>", "", markdown_content, flags=re.IGNORECASE).strip()
    clean_md = re.sub(r"</?think>", "", clean_md, flags=re.IGNORECASE).strip()
    
    lines = clean_md.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        trimmed = line.strip()
        
        if not trimmed:
            i += 1
            continue
            
        # 分割线
        if re.match(r"^(\-{3,}|\*{3,}|_{3,})$", trimmed):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="1" w:color="CBD5E1"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            i += 1
            continue
            
        # 标题 # ## ###
        h_match = re.match(r"^(#{1,4})\s+(.*)$", trimmed)
        if h_match:
            lvl = len(h_match.group(1))
            h_text = h_match.group(2).strip()
            if lvl == 1:
                p = add_styled_paragraph(doc, h_text, font_size=14, bold=True, color_rgb=RGBColor(15, 23, 42), space_before=16, space_after=6)
                p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="2" w:color="CBD5E1"/></w:pBdr>')
                p._p.get_or_add_pPr().append(p_border)
            elif lvl == 2:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.3
                p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="20" w:space="4" w:color="4F46E5"/></w:pBdr>')
                p._p.get_or_add_pPr().append(p_border)
                run = p.add_run(h_text)
                run.bold = True
                run.font.size = Pt(12)
                run.font.color.rgb = RGBColor(30, 27, 75)
                set_run_fonts(run)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(3)
                run_dot = p.add_run("▫️ ")
                run_dot.font.size = Pt(9)
                run = p.add_run(h_text)
                run.bold = True
                run.font.size = Pt(10.5)
                run.font.color.rgb = RGBColor(55, 48, 163)
                set_run_fonts(run)
            i += 1
            continue
            
        # 表格 Table
        if trimmed.startswith("|"):
            table_lines = []
            while i < len(lines) and (lines[i].strip().startswith("|") or (re.match(r"^[\:\-\s|]+$", lines[i].strip()) and lines[i].strip())):
                table_lines.append(lines[i].strip())
                i += 1
                
            if table_lines:
                rows_data = []
                for tl in table_lines:
                    if re.match(r"^\|?[\s\-:|]+\|?$", tl) and ("-" in tl or ":" in tl):
                        continue
                    raw_cells = tl.split("|")
                    if tl.startswith("|"): raw_cells.pop(0)
                    if tl.endswith("|") and raw_cells: raw_cells.pop(-1)
                    cells = [c.strip() for c in raw_cells]
                    if cells:
                        rows_data.append(cells)
                        
                if rows_data:
                    num_cols = max(len(r) for r in rows_data)
                    word_table = doc.add_table(rows=len(rows_data), cols=num_cols)
                    word_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    word_table.autofit = False
                    
                    if num_cols == 6:
                        col_widths = [Inches(1.35), Inches(1.05), Inches(2.00), Inches(0.75), Inches(0.72), Inches(0.90)]
                    elif num_cols == 3:
                        col_widths = [Inches(2.00), Inches(1.80), Inches(2.97)]
                    elif num_cols == 5:
                        col_widths = [Inches(1.10), Inches(1.20), Inches(1.20), Inches(1.20), Inches(2.07)]
                    else:
                        even_w = 6.77 / num_cols
                        col_widths = [Inches(even_w) for _ in range(num_cols)]
                    
                    for r_idx, row_cells in enumerate(rows_data):
                        is_hdr = (r_idx == 0)
                        for c_idx in range(num_cols):
                            cell = word_table.cell(r_idx, c_idx)
                            if c_idx < len(col_widths):
                                cell.width = col_widths[c_idx]
                            set_cell_margins(cell, top=90, bottom=90, left=110, right=110)
                            set_cell_borders(cell, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1")
                            
                            if is_hdr:
                                set_cell_shading(cell, "F1F5F9")
                            elif r_idx % 2 == 1:
                                set_cell_shading(cell, "FAFAFA")
                                
                            cell_text = row_cells[c_idx] if c_idx < len(row_cells) else ""
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            p.paragraph_format.line_spacing = 1.15
                            
                            is_mono = False
                            is_numeric = bool(re.match(r"^[\+\-]?[\d\.]+%?$", cell_text))
                            if is_numeric:
                                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                            elif c_idx == 1 and ("alpha" in cell_text.lower() or "momentum" in cell_text.lower() or "bias" in cell_text.lower()):
                                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                is_mono = True
                                
                            clean_text = cell_text.replace("**", "").replace("`", "")
                            run = p.add_run(clean_text)
                            run.font.size = Pt(8.5 if not is_hdr else 9)
                            
                            if is_hdr:
                                run.bold = True
                                run.font.color.rgb = RGBColor(15, 23, 42)
                                set_run_fonts(run)
                            else:
                                if "多头" in clean_text or "积极" in clean_text:
                                    run.bold = True
                                    run.font.color.rgb = RGBColor(185, 28, 28)
                                    set_run_fonts(run)
                                elif "空头" in clean_text or "谨慎" in clean_text:
                                    run.bold = True
                                    run.font.color.rgb = RGBColor(4, 120, 87)
                                    set_run_fonts(run)
                                elif is_mono or (c_idx == 2 and ("corr(" in clean_text or "rank(" in clean_text or "close" in clean_text)):
                                    run.font.size = Pt(8.0)
                                    run.font.color.rgb = RGBColor(79, 70, 229)
                                    set_run_fonts(run, is_code=True)
                                else:
                                    run.font.color.rgb = RGBColor(51, 65, 85)
                                    set_run_fonts(run)
                    doc.add_paragraph() # 表格后空行
            continue
            
        # 引用与 Callout (> ...)
        if trimmed.startswith(">") or trimmed.startswith("&gt;"):
            quote_lines = []
            while i < len(lines) and (lines[i].strip().startswith(">") or lines[i].strip().startswith("&gt;")):
                ql = re.sub(r"^(&gt;|>)\s?", "", lines[i].strip())
                quote_lines.append(ql)
                i += 1
                
            expanded = []
            for l in quote_lines:
                s = re.sub(r"<br\s*/?>", "\n", l, flags=re.IGNORECASE)
                s = re.sub(r"([。！？；])\s*(\d+\.\s+)", r"\1\n\2", s)
                for item in s.split("\n"):
                    if item.strip():
                        expanded.append(item.strip())
                        
            first_l = expanded[0] if expanded else ""
            is_sug = ("💡" in first_l or "建议" in first_l or "策略" in first_l)
            is_warn = not is_sug and ("⚠️" in first_l or "风控" in first_l or "风险" in first_l or "注意" in first_l)
            is_quest = not is_sug and not is_warn and ("💬" in first_l or "追问" in first_l or "问：" in first_l or "问题" in first_l)
            
            if is_sug:
                c_type = "suggestion"
                clean_title = "操作建议"
            elif is_warn:
                c_type = "warning"
                clean_title = "风控警示"
            elif is_quest:
                c_type = "question"
                clean_title = "投资者追问"
            else:
                c_type = "default"
                clean_title = "重点提示"
                
            items = []
            for l_idx, l_str in enumerate(expanded):
                clean_l = re.sub(r"^[\s💡⚠️💬📌🔔⭐]+", "", l_str).strip()
                if l_idx == 0 and ("建议" in clean_l or "风控" in clean_l or "警示" in clean_l or "追问" in clean_l):
                    clean_title = clean_l.replace("**", "").replace(":", "").replace("：", "")
                else:
                    items.append(clean_l)
            if not items and expanded:
                items = [clean_title]
                clean_title = "操作建议" if is_sug else ("风控警示" if is_warn else ("投资者追问" if is_quest else "重点提示"))
                
            add_callout_box(doc, clean_title, items, callout_type=c_type)
            continue
            
        # 列表项
        if re.match(r"^[\-\*]\s+", trimmed) or re.match(r"^\d+\.\s+", trimmed):
            add_styled_paragraph(doc, trimmed, font_size=9.5, space_before=2, space_after=2, line_spacing=1.3)
            i += 1
            continue
            
        # 常规段落
        p_lines = []
        while i < len(lines):
            c_trim = lines[i].strip()
            if not c_trim or c_trim.startswith("#") or c_trim.startswith("|") or c_trim.startswith(">") or c_trim.startswith("&gt;") or re.match(r"^[\-\*]\s+", c_trim) or re.match(r"^\d+\.\s+", c_trim):
                break
            p_lines.append(c_trim)
            i += 1
            
        if p_lines:
            add_styled_paragraph(doc, " ".join(p_lines), font_size=10, space_before=4, space_after=4, line_spacing=1.4)
        else:
            add_styled_paragraph(doc, trimmed, font_size=10, space_before=4, space_after=4, line_spacing=1.4)
            i += 1

    # 6. 底部权威量化免责声明
    p_disc = doc.add_paragraph()
    p_disc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_disc.paragraph_format.space_before = Pt(28)
    p_disc.paragraph_format.space_after = Pt(0)
    run_disc = p_disc.add_run("【免责声明】本报告基于量化多因子截面模型与大模型自主推理生成，仅供专业投研分析与决策参考，不构成任何实质性投资建议或买卖依据。市场有风险，投资需谨慎。\nAshare Vibe AI Quantitative Research System · All Rights Reserved")
    run_disc.font.size = Pt(8)
    run_disc.font.color.rgb = RGBColor(148, 163, 184)
    set_run_fonts(run_disc)
    
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
